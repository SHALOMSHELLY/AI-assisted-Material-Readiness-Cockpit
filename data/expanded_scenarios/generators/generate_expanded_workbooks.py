"""CLI that writes expanded workbooks. It is intentionally not run by setup."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

if __package__ in {None, ""}:
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[3]))
    from data.expanded_scenarios.generators import ALL_SCENARIOS  # type: ignore
    from data.expanded_scenarios.generators.schemas import SHEET_COLUMNS  # type: ignore
else:
    from . import ALL_SCENARIOS
    from .schemas import SHEET_COLUMNS

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "generated"
FORBIDDEN_OUTPUT = Path(__file__).resolve().parents[3] / "test_data" / "workbooks"


def _selected(scenario_id: str | None, category: str | None) -> Iterable[object]:
    scenarios = list(ALL_SCENARIOS.values())
    if scenario_id:
        if scenario_id not in ALL_SCENARIOS:
            raise SystemExit(f"Unknown scenario: {scenario_id}")
        scenarios = [ALL_SCENARIOS[scenario_id]]
    if category:
        scenarios = [scenario for scenario in scenarios if scenario.category == category]
    return sorted(scenarios, key=lambda scenario: scenario.scenario_id)


def _safe_output(path: Path) -> Path:
    resolved = path.resolve()
    forbidden = FORBIDDEN_OUTPUT.resolve()
    if resolved == forbidden or forbidden in resolved.parents:
        raise SystemExit("Refusing to write into the legacy test_data/workbooks directory.")
    return resolved


def write_workbook(path: Path, scenario: object, overwrite: bool) -> None:
    """Write typed values and stable formatting without replacing by default."""

    if path.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite existing workbook: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(); workbook.remove(workbook.active)
    sheets = {name: rows for name, rows in scenario.sheets.items() if name not in scenario.omitted_sheets}
    sheets.update(scenario.extra_sheets or {})
    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        columns = SHEET_COLUMNS.get(sheet_name, list(rows[0]) if rows else ["Note"])
        worksheet.append(columns)
        for row in rows:
            worksheet.append([row.get(column) for column in columns])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for cells in worksheet.columns:
            width = min(32, max(12, max(len(str(cell.value or "")) for cell in cells) + 2))
            worksheet.column_dimensions[cells[0].column_letter].width = width
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--scenario", choices=sorted(ALL_SCENARIOS))
    parser.add_argument("--category", choices=("small", "medium", "large"))
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if args.scenario and args.category:
        parser.error("--scenario and --category are mutually exclusive")
    output = _safe_output(args.output)
    selected = list(_selected(args.scenario, args.category))
    for scenario in selected:
        path = output / scenario.category / f"{scenario.scenario_id}.xlsx"
        counts = ", ".join(f"{name}={count}" for name, count in scenario.row_counts.items())
        if args.dry_run:
            print(f"{scenario.scenario_id}: {path} ({counts})")
        else:
            write_workbook(path, scenario, args.overwrite)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
