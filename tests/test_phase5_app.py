"""Phase 5 UI, state, and export tests."""

import csv
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from src.export_results import results_to_csv, results_to_excel
from src.shortage_engine import run_shortage_analysis
from src.ui_state import upload_signature
from src.workbook_reader import read_workbook


ROOT = Path(__file__).resolve().parents[1]
TEST_DATE = date(2026, 9, 25)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def start_app() -> AppTest:
    """Start the app with the fixed test date."""

    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=10).run()
    app.date_input[0].set_value(TEST_DATE).run()
    return app


def upload_and_validate(path: Path) -> AppTest:
    """Upload one workbook and click validation."""

    app = start_app()
    app.file_uploader[0].upload(path.name, path.read_bytes(), XLSX_MIME).run()
    next(button for button in app.button if button.label.startswith("Validate Workbook")).click().run()
    return app


def metric_value(app: AppTest, english_label: str) -> str:
    return next(metric.value for metric in app.metric if metric.label.startswith(english_label))


def test_app_starts_with_today_as_default():
    app = AppTest.from_file(str(ROOT / "app.py"), default_timeout=10).run()
    assert not app.exception
    assert app.date_input[0].value == date.today()
    assert any(title.value == "Material Readiness Cockpit" for title in app.title)
    assert any("Deterministic material readiness" in item.value for item in app.markdown)
    assert not any(button.label == "Load Demonstration Data" for button in app.button)


def test_upload_requires_explicit_validation_before_analysis():
    app = start_app()
    path = ROOT / "test_data" / "sample_workbook.xlsx"
    app.file_uploader[0].upload(path.name, path.read_bytes(), XLSX_MIME).run()
    assert any("Workbook loaded" in item.value for item in app.success)
    assert next(button for button in app.button if button.label == "Run Analysis").disabled
    next(button for button in app.button if button.label.startswith("Validate Workbook")).click().run()
    assert not next(button for button in app.button if button.label == "Run Analysis").disabled


def test_valid_sample_upload_validation_and_analysis():
    app = upload_and_validate(ROOT / "test_data" / "sample_workbook.xlsx")
    assert metric_value(app, "Blocking Problems") == "0"
    assert metric_value(app, "Check Notes") == "0"
    assert metric_value(app, "Excluded Rows") == "0"
    next(button for button in app.button if button.label == "Run Analysis").click().run()
    assert not app.exception
    assert metric_value(app, "Requirement Quantity") == "16 EA"
    assert metric_value(app, "Inventory Allocated") == "10"
    assert metric_value(app, "Incoming PO Allocated") == "6"
    assert metric_value(app, "Uncovered Quantity") == "0"
    assert any("16 - (10 + 6) = 0 EA" in code.value for code in app.code)
    assert any("Low" in text.value and "P4" in text.value for text in app.markdown)


def test_invalid_workbook_shows_errors_and_cannot_run_analysis():
    app = upload_and_validate(ROOT / "test_data" / "invalid_validation_example.xlsx")
    assert int(metric_value(app, "Blocking Problems")) > 0
    assert next(button for button in app.button if button.label == "Run Analysis").disabled
    assert not app.exception


def test_analysis_date_change_clears_prior_analysis():
    app = upload_and_validate(ROOT / "test_data" / "sample_workbook.xlsx")
    next(button for button in app.button if button.label == "Run Analysis").click().run()
    assert app.code
    app.date_input[0].set_value(date(2026, 9, 26)).run()
    assert not app.code
    assert not app.metric
    assert next(button for button in app.button if button.label == "Run Analysis").disabled


def test_signature_distinguishes_same_name_and_size_content_and_date():
    first = b"same-size-A"
    second = b"same-size-B"
    assert len(first) == len(second)
    assert upload_signature(first, TEST_DATE) != upload_signature(second, TEST_DATE)
    assert upload_signature(first, TEST_DATE) != upload_signature(first, date(2026, 9, 26))


def test_excel_and_csv_exports_contain_real_sample_values():
    result = run_shortage_analysis(
        read_workbook(ROOT / "test_data" / "sample_workbook.xlsx"), TEST_DATE
    )
    workbook = load_workbook(BytesIO(results_to_excel(result)), read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "Shortage_Results", "Material_Requirements", "Validation_Summary", "Validation_Issues"
    ]
    sheet = workbook["Shortage_Results"]
    headers = [cell.value for cell in sheet[1]]
    values = dict(zip(headers, [cell.value for cell in sheet[2]]))
    assert values["Required_Quantity"] == 16
    assert values["Inventory_Allocated"] == 10
    assert values["Incoming_PO_Allocated"] == 6
    assert values["Shortage_Quantity"] == 0
    assert values["Priority"] == "Low"
    assert values["Priority_Rule_ID"] == "P4"

    csv_text = results_to_csv(result).decode("utf-8-sig")
    csv_row = next(csv.DictReader(StringIO(csv_text)))
    assert csv_row["Required_Quantity"] == "16.0"
    assert csv_row["Inventory_Allocated"] == "10.0"
    assert csv_row["Incoming_PO_Allocated"] == "6.0"
    assert csv_row["Shortage_Quantity"] == "0.0"
    assert csv_row["Priority"] == "Low"
    assert csv_row["Priority_Rule_ID"] == "P4"


def test_validation_issues_sheet_exports_warning_detail():
    sheets = read_workbook(ROOT / "test_data" / "sample_workbook.xlsx")
    sheets["Incoming_PO"].loc[0, "Expected_Receipt_Date"] = date(2026, 9, 24)
    result = run_shortage_analysis(sheets, TEST_DATE)
    workbook = load_workbook(BytesIO(results_to_excel(result)), read_only=True, data_only=True)
    issue_sheet = workbook["Validation_Issues"]
    headers = [cell.value for cell in issue_sheet[1]]
    rows = [dict(zip(headers, values)) for values in issue_sheet.iter_rows(min_row=2, values_only=True)]
    assert any(row["code"] == "OVERDUE_INCOMING_PO" and row["severity"] == "WARNING" for row in rows)


def test_empty_analysis_still_offers_downloads_and_validation_summary():
    sheets = read_workbook(ROOT / "test_data" / "sample_workbook.xlsx")
    sheets["Production_Orders"].loc[0, "Order_Status"] = "COMPLETED"
    result = run_shortage_analysis(sheets, TEST_DATE)
    assert result.shortage_results.empty
    workbook = load_workbook(BytesIO(results_to_excel(result)), read_only=True)
    assert workbook["Validation_Summary"]["A2"].value is True

    source = load_workbook(ROOT / "test_data" / "sample_workbook.xlsx")
    production = source["Production_Orders"]
    status_column = [cell.value for cell in production[1]].index("Order_Status") + 1
    production.cell(row=2, column=status_column, value="COMPLETED")
    content = BytesIO()
    source.save(content)
    app = start_app()
    app.file_uploader[0].upload("no_active_demand.xlsx", content.getvalue(), XLSX_MIME).run()
    next(button for button in app.button if button.label.startswith("Validate Workbook")).click().run()
    next(button for button in app.button if button.label == "Run Analysis").click().run()
    labels = [button.label for button in app.download_button]
    assert "Download Shortage Results · Excel" in labels
    assert "Download Shortage Results · CSV" in labels
    assert "Download Exception Register · CSV" in labels
    assert "Download Order Summary · CSV" in labels
